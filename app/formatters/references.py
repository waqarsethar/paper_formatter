"""Reference list formatter: parses and reformats the bibliography section."""

from __future__ import annotations

import re

from docx.shared import Pt, Inches

from app.core.document_parser import find_section_by_heading, is_reference_heading


# ---------------------------------------------------------------------------
# Common reference-section heading names (passed to find_section_by_heading)
# ---------------------------------------------------------------------------

_REF_HEADINGS = [
    "References",
    "Bibliography",
    "Works Cited",
    "Literature Cited",
    "Citations",
    "Reference List",
    "Cited Literature",
    "Literature References",
]


# ---------------------------------------------------------------------------
# Regex helpers for parsing individual references
# ---------------------------------------------------------------------------

# DOI anywhere in the string.
_DOI_RE = re.compile(r'(?:doi:\s*|https?://doi\.org/)(\S+)', re.IGNORECASE)

# Year in parentheses — typical of APA style: Author (2024).
_YEAR_PAREN_RE = re.compile(r'\((\d{4}[a-z]?)\)')

# Year followed by a period — typical of Vancouver/numbered styles.
_YEAR_PLAIN_RE = re.compile(r'(\d{4}[a-z]?)\.')

# Volume(issue), pages pattern: e.g. 12(3), 45-67 or 12, 45-67
_VOL_ISSUE_PAGES_RE = re.compile(
    r'(\d+)\s*\((\d+)\)\s*[,:]\s*([\d\-–]+)'
)

# Volume and pages without explicit issue: e.g. 12, 45-67
_VOL_PAGES_RE = re.compile(
    r'(\d+)\s*[,:]\s*([\d\-–]+)'
)

# Numbered reference prefix: "1. ", "1) ", "[1] "
_NUMBER_PREFIX_RE = re.compile(r'^\s*(?:\[?\d+[\].)]\s*)')


# ---------------------------------------------------------------------------
# Section finder
# ---------------------------------------------------------------------------

def find_references_section(doc) -> tuple[int, int] | None:
    """Locate the references / bibliography section in the document.

    Returns a ``(start_idx, end_idx)`` tuple of paragraph indices, or
    ``None`` if no matching heading is found.  *start_idx* is the heading
    paragraph itself; the actual reference entries start at
    ``start_idx + 1``.
    """
    return find_section_by_heading(doc, _REF_HEADINGS)


# ---------------------------------------------------------------------------
# Reference parser
# ---------------------------------------------------------------------------

def parse_reference(text: str) -> dict | None:
    """Best-effort parse of a reference string into structured fields.

    Attempts to extract:
        authors, year, title, journal, volume, issue, pages, doi

    Returns ``None`` when fewer than 3 fields can be identified (low
    confidence).
    """
    fields: dict[str, str] = {}
    remaining = text.strip()

    # Strip leading number prefix (e.g. "1. " or "[1] ").
    remaining = _NUMBER_PREFIX_RE.sub("", remaining).strip()

    # --- DOI ---
    doi_match = _DOI_RE.search(remaining)
    if doi_match:
        fields["doi"] = doi_match.group(1).rstrip(".")
        remaining = remaining[: doi_match.start()] + remaining[doi_match.end() :]
        remaining = remaining.strip().rstrip(".")

    # --- Year ---
    year_match = _YEAR_PAREN_RE.search(remaining)
    if year_match:
        fields["year"] = year_match.group(1)
        year_pos = year_match.start()
        before_year = remaining[:year_pos].strip().rstrip(".,")
        after_year = remaining[year_match.end():].strip().lstrip(".,").strip()
    else:
        year_match = _YEAR_PLAIN_RE.search(remaining)
        if year_match:
            fields["year"] = year_match.group(1)
            year_pos = year_match.start()
            before_year = remaining[:year_pos].strip().rstrip(".,")
            after_year = remaining[year_match.end():].strip()
        else:
            before_year = remaining
            after_year = ""

    # --- Authors (text before the year) ---
    if before_year:
        fields["authors"] = before_year.strip()

    # --- Volume, issue, pages ---
    vim = _VOL_ISSUE_PAGES_RE.search(after_year)
    if vim:
        fields["volume"] = vim.group(1)
        fields["issue"] = vim.group(2)
        fields["pages"] = vim.group(3)
        # Text before vol/issue is typically "Title. Journal"
        pre_vol = after_year[: vim.start()].strip().rstrip(",. ")
        post_vol = after_year[vim.end():].strip().lstrip(",. ")
    else:
        vp = _VOL_PAGES_RE.search(after_year)
        if vp:
            fields["volume"] = vp.group(1)
            fields["pages"] = vp.group(2)
            pre_vol = after_year[: vp.start()].strip().rstrip(",. ")
            post_vol = after_year[vp.end():].strip().lstrip(",. ")
        else:
            pre_vol = after_year
            post_vol = ""

    # --- Title and journal (heuristic split on period) ---
    if pre_vol:
        # Split on first period that is likely a sentence boundary
        # (not initials like "J." or "U.S.").
        parts = re.split(r'(?<=[a-z\?\!])\.\s+', pre_vol, maxsplit=1)
        if len(parts) == 2:
            fields["title"] = parts[0].strip().rstrip(".")
            fields["journal"] = parts[1].strip().rstrip(",. ")
        elif len(parts) == 1:
            # Only one chunk: try splitting on the last comma.
            comma_parts = pre_vol.rsplit(",", 1)
            if len(comma_parts) == 2 and len(comma_parts[1].strip()) > 2:
                fields["title"] = comma_parts[0].strip().rstrip(".")
                fields["journal"] = comma_parts[1].strip().rstrip(",. ")
            else:
                fields["title"] = pre_vol.strip().rstrip(".")

    # Confidence check: require at least 3 extracted fields.
    if len(fields) < 3:
        return None

    return fields


# ---------------------------------------------------------------------------
# Reference formatter
# ---------------------------------------------------------------------------

def format_reference(fields: dict, template: str, number: int) -> str:
    """Apply *template* to *fields*, substituting placeholders.

    Supported placeholders: ``{num}``, ``{authors}``, ``{year}``,
    ``{title}``, ``{journal}``, ``{volume}``, ``{issue}``, ``{pages}``,
    ``{doi}``.

    Missing fields are replaced with empty strings and extraneous
    punctuation around them is cleaned up.
    """
    replacements = {
        "num": str(number),
        "authors": fields.get("authors", ""),
        "year": fields.get("year", ""),
        "title": fields.get("title", ""),
        "journal": fields.get("journal", ""),
        "volume": fields.get("volume", ""),
        "issue": fields.get("issue", ""),
        "pages": fields.get("pages", ""),
        "doi": fields.get("doi", ""),
    }

    result = template
    for key, value in replacements.items():
        result = result.replace("{" + key + "}", value)

    # Clean up artefacts from empty fields: collapse multiple commas /
    # periods, empty parentheses, leading/trailing whitespace.
    result = re.sub(r'\(\)', '', result)           # empty parens
    result = re.sub(r',\s*,', ',', result)         # double commas
    result = re.sub(r'\.\s*\.', '.', result)       # double periods
    result = re.sub(r',\s*\.', '.', result)        # comma before period
    result = re.sub(r'\s{2,}', ' ', result)        # collapse whitespace
    result = result.strip().rstrip(",").strip()

    return result


# ---------------------------------------------------------------------------
# Main entry-point
# ---------------------------------------------------------------------------

def apply_references(doc, config: dict) -> dict:
    """Parse and reformat the reference list according to journal style.

    Args:
        doc: A ``python-docx`` ``Document`` instance.
        config: Journal configuration dict.  Must contain
            ``config["reference_style"]`` with keys ``numbering``,
            ``format``, ``hanging_indent``, and ``font_size``.

    Returns:
        A dict with ``"warnings"`` (list[str]) and ``"stats"`` (dict).
    """
    warnings: list[str] = []
    stats: dict = {"references_found": 0, "references_reformatted": 0}

    ref_config = config.get("reference_style", {})
    numbering = ref_config.get("numbering", "unnumbered")
    template = ref_config.get(
        "format",
        "{authors} ({year}). {title}. {journal}, {volume}({issue}), {pages}. {doi}",
    )
    hanging_indent = ref_config.get("hanging_indent", 0.5)
    font_size = ref_config.get("font_size", 10.0)
    font_size_pt = Pt(font_size)

    # Step 1 — find the references section.
    section_range = find_references_section(doc)
    if section_range is None:
        warnings.append(
            "Could not locate a References / Bibliography section heading. "
            "No reference reformatting applied."
        )
        return {"warnings": warnings, "stats": stats}

    ref_start, ref_end = section_range
    # Actual entries begin one paragraph after the heading.
    entry_start = ref_start + 1

    paragraphs = doc.paragraphs
    references_found = 0
    references_reformatted = 0
    number = 1

    for idx in range(entry_start, ref_end):
        para = paragraphs[idx]
        text = para.text.strip()
        if not text:
            continue  # skip blank lines

        references_found += 1

        # Step 2 — parse.
        fields = parse_reference(text)
        if fields is None:
            warnings.append(
                f"Could not parse reference (entry #{references_found}): "
                f"'{text[:80]}{'...' if len(text) > 80 else ''}'. Left unchanged."
            )
            # Still apply layout formatting even if we can't parse content.
            _apply_paragraph_formatting(
                para, hanging_indent, font_size_pt
            )
            number += 1
            continue

        # Step 3 — reformat.
        new_text = format_reference(fields, template, number)

        # If numbering is "numbered" and the template doesn't already
        # contain {num}, prepend the number.
        if numbering == "numbered" and "{num}" not in template:
            new_text = f"{number}. {new_text}"

        # Replace paragraph text while preserving a single run.
        _replace_paragraph_text(para, new_text)

        # Step 4 — apply formatting.
        _apply_paragraph_formatting(para, hanging_indent, font_size_pt)

        references_reformatted += 1
        number += 1

    stats["references_found"] = references_found
    stats["references_reformatted"] = references_reformatted

    return {"warnings": warnings, "stats": stats}


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _replace_paragraph_text(paragraph, new_text: str) -> None:
    """Replace all text in *paragraph* with *new_text*.

    Clears every existing run and writes *new_text* into the first run,
    preserving its character formatting.
    """
    runs = paragraph.runs
    if not runs:
        # No runs — add one.
        paragraph.add_run(new_text)
        return

    # Preserve formatting from the first run.
    first_run = runs[0]
    first_run.text = new_text

    # Clear remaining runs.
    for run in runs[1:]:
        run.text = ""


def _apply_paragraph_formatting(paragraph, hanging_indent: float, font_size_pt) -> None:
    """Apply hanging indent and font size to a reference paragraph."""
    pf = paragraph.paragraph_format
    indent_val = Inches(hanging_indent)
    pf.left_indent = indent_val
    pf.first_line_indent = -indent_val

    for run in paragraph.runs:
        run.font.size = font_size_pt
